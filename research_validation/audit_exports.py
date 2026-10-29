"""Run the unchanged production exporter in an isolated copy and read all data back."""
import contextlib
import hashlib
import io
import json
import math
import runpy
import shutil
import sys
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
# optional local dependency folder (not part of the repository)
sys.path.insert(0, str(ROOT / ".paper_work/export_deps"))
import pandas as pd
import pyreadstat
import openpyxl

WORK = ROOT / "research_validation/export_run"
HERE = WORK / "dataset"
HERE.mkdir(parents=True, exist_ok=True)
for name in ("build_dataset.py", "lookups.py", "osym_ozet.py"):
    shutil.copy2(ROOT / "dataset" / name, HERE / name)
shutil.copy2(ROOT / "data.json", WORK / "data.json")
sys.path.insert(0, str(HERE))
log = io.StringIO()
with contextlib.redirect_stdout(log):
    env = runpy.run_path(str(HERE / "build_dataset.py"))
(WORK / "build_log.txt").write_text(log.getvalue(), encoding="utf-8")
OUT = HERE / "out"
report = {"timestamp_utc": datetime.now(timezone.utc).isoformat(),
          "packages": {"pandas": pd.__version__, "pyreadstat": pyreadstat.__version__, "openpyxl": openpyxl.__version__},
          "source_sha256": {n: hashlib.sha256((ROOT / n).read_bytes()).hexdigest()
                            for n in ("data.json", "dataset/build_dataset.py")},
          "checks": {}, "errors": [],
          "scope": "Round-trip of all 7170x33 data cells. XLSX formulas inspected separately; no native SPSS application test."}


def missing(value):
    return value is None or bool(pd.isna(value))


def compare(name, actual, expected):
    issues = []
    if list(actual.columns) != list(expected.columns) or actual.shape != expected.shape:
        issues.append("shape/columns")
    else:
        for col in expected.columns:
            for i, (a, e) in enumerate(zip(actual[col], expected[col])):
                same = missing(a) and missing(e)
                if not missing(a) and not missing(e):
                    same = (math.isclose(float(a), float(e), rel_tol=0, abs_tol=1e-10)
                            if isinstance(e, (int, float)) else a == e)
                if not same:
                    issues.append({"row": i + 2, "column": col, "actual": str(a), "expected": str(e)})
    report["checks"][name] = {"shape": list(actual.shape), "cells": int(expected.size), "mismatches": len(issues), "examples": issues[:10]}
    if issues:
        report["errors"].append(name)


# Previous independent audit snapshot anchors the production table before serialization.
baseline = pd.DataFrame(json.loads((ROOT / "research_validation/dataset_records.json").read_text(encoding="utf-8")))
compare("coded_table_vs_previous_audit", env["coded"], baseline[env["COLS"]])
strings = {c: str for c in env["COLS"] if not pd.api.types.is_numeric_dtype(env["labeled"][c])}
csv = pd.read_csv(OUT / "ozyes_veriseti.csv", dtype=strings, encoding="utf-8-sig")
compare("csv", csv, env["labeled"])
sav, meta = pyreadstat.read_sav(OUT / "ozyes_veriseti.sav", apply_value_formats=False)
compare("sav_numeric", sav, env["coded"])
sav_labels, _ = pyreadstat.read_sav(OUT / "ozyes_veriseti.sav", apply_value_formats=True, formats_as_category=False)
compare("sav_labeled", sav_labels, env["labeled"])
label_ok = meta.column_names_to_labels == {d[0]: d[1] for d in env["DICT"]}
value_ok = meta.variable_value_labels == {c: {float(k): v for k, v in labels.items()} for c, labels in env["VALUE_LABELS"].items()}
report["checks"]["sav_metadata"] = {"column_labels_equal": label_ok, "value_labels_equal": value_ok, "measure": meta.variable_measure}
if not label_ok or not value_ok:
    report["errors"].append("sav_metadata")
wb = openpyxl.load_workbook(OUT / "ozyes_veriseti.xlsx", data_only=False)
rows = list(wb["Veri"].values)
compare("xlsx_data", pd.DataFrame(rows[1:], columns=rows[0]), env["labeled"])
formulas = {s.title: {c.coordinate: c.value for row in s for c in row if c.data_type == "f"} for s in wb}
report["xlsx"] = {"sheets": wb.sheetnames, "formulas_per_sheet": {s: len(f) for s, f in formulas.items()},
                  "program_panel_rows": wb["Program_Panel"].max_row - 1,
                  "university_rows": wb["Üniversiteler"].max_row - 1,
                  "native_excel_recalculation": "see excel_recalculation.json when present"}
(WORK / "formulas.json").write_text(json.dumps(formulas, ensure_ascii=False, indent=2), encoding="utf-8")
report["output_sha256"] = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in OUT.iterdir()}
(ROOT / "research_validation/export_audit.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps({"checks": report["checks"], "xlsx": report["xlsx"], "errors": report["errors"]}, ensure_ascii=False, indent=2))
raise SystemExit(bool(report["errors"]))
